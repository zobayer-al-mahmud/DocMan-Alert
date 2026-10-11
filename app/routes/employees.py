import os
import secrets
from datetime import date

from flask import (Blueprint, render_template, redirect, url_for, flash, current_app,
                   send_from_directory, request, abort)
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from sqlalchemy.orm import joinedload

from app import db
from app.models import (
    Department, Employee, EmployeeCompliance, EmployeeComplianceVersion, Factory,
    HealthCheckup, HealthCheckupVersion, OcrResult, User,
)
from app.forms import EmployeeForm, EmployeeComplianceForm, HealthCheckupForm
from app.utils import (
    roles_required, compute_status, ALL_STATUSES, log_event, valid_email,
    expiry_before_issue, matching_alert_users,
)

employees_bp = Blueprint("employees", __name__, url_prefix="/employees")

MANAGE_ROLES = (
    "Super Admin", "Administration", "Factory Admin",
    "Compliance Manager", "Compliance User", "Transport Manager", "Transport User",
)
DELETE_ROLES = (
    "Super Admin", "Administration", "Factory Admin", "Compliance Manager", "Transport Manager",
)

LICENSE_TYPE_SUGGESTIONS = ["Electrician License", "Medical License", "Boiler Operator License", "Other"]


# --- Scope helpers ----------------------------------------------------------

def get_accessible_factories():
    if current_user.can_access_all_factories():
        return Factory.query.order_by(Factory.name).all()
    if current_user.factory_id:
        return Factory.query.filter_by(id=current_user.factory_id).order_by(Factory.name).all()
    return []


def can_access_factory(factory_id):
    if current_user.can_access_all_factories():
        return True
    return current_user.factory_id == factory_id


def apply_employee_scope(query):
    if current_user.can_access_all_factories():
        return query
    if current_user.factory_id:
        return query.filter(Employee.factory_id == current_user.factory_id)
    return query.filter(db.false())


def _get_employee_or_403(employee_id):
    employee = Employee.query.get_or_404(employee_id)
    if not can_access_factory(employee.factory_id):
        abort(403)
    return employee


def _save_employee_file(file_storage, kind, employee, record_id, version_no):
    """Same extension rules as factory documents. The stored name includes the version."""
    original = secure_filename(file_storage.filename or "")
    if not original or "." not in original:
        raise ValueError("Invalid file name.")
    ext = original.rsplit(".", 1)[1].lower()
    if ext not in current_app.config["ALLOWED_EXTENSIONS"]:
        raise ValueError("File type is not allowed.")
    subfolder = f"employees/factory_{employee.factory_id}"
    unique_name = f"{kind}{record_id}_v{version_no}_{original}"
    relative_path = f"{subfolder}/{unique_name}"
    folder_path = os.path.join(current_app.config["UPLOAD_FOLDER"], subfolder)
    os.makedirs(folder_path, exist_ok=True)
    file_storage.stream.seek(0)
    file_storage.save(os.path.join(current_app.config["UPLOAD_FOLDER"], relative_path))
    return relative_path


def _set_factory_choices(form):
    factories = get_accessible_factories()
    form.factory_id.choices = [(f.id, f.name) for f in factories]
    if hasattr(form, "department_id"):
        departments = Department.query.order_by(Department.name).all()
        form.department_id.choices = [(0, "-- None --")] + [(d.id, d.name) for d in departments]


def _department_for_factory(department_id, factory_id):
    """factory_id is unused. A department can be used at any factory."""
    if not department_id:
        return None
    department = db.session.get(Department, department_id)
    if not department:
        return False
    return department


def _set_record_recipient_choices(form, factory_id, department_id):
    """Users in this employee's factory and department."""
    users = matching_alert_users(factory_id, department_id)
    eligible = [user for user in users if valid_email(user.email)]
    ineligible = [user for user in users if not valid_email(user.email)]
    form.alert_recipients.choices = [
        (user.id, f"{user.full_name or user.username} — {user.email}") for user in eligible
    ]
    return eligible, ineligible, {user.id for user in eligible}


def _apply_record_recipients(record, employee, form, eligible_ids):
    """Keep only selected users who belong to this employee's factory and department."""
    allowed = {
        user.id for user in matching_alert_users(employee.factory_id, employee.department_id)
    }
    selected = [
        uid for uid in (form.alert_recipients.data or [])
        if uid in eligible_ids and uid in allowed
    ]
    record.alert_recipients = User.query.filter(User.id.in_(selected)).all() if selected else []


def _attach_license_file(record, employee, file_storage):
    version_no = (max((v.version_no for v in record.file_versions), default=0)) + 1
    path = _save_employee_file(file_storage, "lic", employee, record.id, version_no)
    record.file_path = path
    db.session.add(EmployeeComplianceVersion(
        compliance_id=record.id,
        version_no=version_no,
        file_path=path,
        uploaded_by=current_user.id,
    ))


def _attach_checkup_file(record, employee, file_storage):
    version_no = (max((v.version_no for v in record.file_versions), default=0)) + 1
    path = _save_employee_file(file_storage, "chk", employee, record.id, version_no)
    record.file_path = path
    db.session.add(HealthCheckupVersion(
        checkup_id=record.id,
        version_no=version_no,
        file_path=path,
        uploaded_by=current_user.id,
    ))


# --- Employees --------------------------------------------------------------

@employees_bp.route("/")
@login_required
def list_employees():
    query = Employee.query.options(joinedload(Employee.factory))
    query = apply_employee_scope(query)

    factory_id = request.args.get("factory_id", type=int)
    section = request.args.get("section", "").strip()
    designation = request.args.get("designation", "").strip()
    search = request.args.get("search", "").strip()
    show_inactive = request.args.get("inactive") == "1"

    if not show_inactive:
        query = query.filter(Employee.is_active.is_(True))
    if factory_id:
        query = query.filter(Employee.factory_id == factory_id)
    if section:
        query = query.filter(Employee.section.ilike(f"%{section}%"))
    if designation:
        query = query.filter(Employee.designation.ilike(f"%{designation}%"))
    if search:
        like = f"%{search}%"
        query = query.filter(db.or_(Employee.name.ilike(like), Employee.employee_code.ilike(like)))

    page = request.args.get("page", 1, type=int)
    pagination = query.order_by(Employee.name).paginate(page=page, per_page=15, error_out=False)

    return render_template(
        "employees/list.html",
        employees=pagination.items,
        pagination=pagination,
        factories=get_accessible_factories(),
        selected_factory_id=factory_id,
        selected_section=section,
        selected_designation=designation,
        search=search,
        show_inactive=show_inactive,
        can_manage=current_user.role.name in MANAGE_ROLES,
    )


@employees_bp.route("/add", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def add_employee():
    form = EmployeeForm()
    _set_factory_choices(form)
    if form.validate_on_submit():
        if not can_access_factory(form.factory_id.data):
            abort(403)
        existing = Employee.query.filter_by(
            factory_id=form.factory_id.data, employee_code=form.employee_code.data.strip()
        ).first()
        if existing:
            flash("An employee with this ID already exists in this factory.", "danger")
            return render_template("employees/form.html", form=form, title="Add Employee")
        department = _department_for_factory(form.department_id.data, form.factory_id.data)
        if department is False:
            flash("Choose a department from the list.", "danger")
            return render_template("employees/form.html", form=form, title="Add Employee")
        receive_alerts = bool(form.receive_alerts.data) and valid_email(form.email.data)
        if form.receive_alerts.data and not receive_alerts:
            flash("Employee alerts stay off until a valid email is saved.", "warning")
        employee = Employee(
            employee_code=form.employee_code.data.strip(),
            name=form.name.data.strip(),
            designation=form.designation.data,
            section=form.section.data,
            factory_id=form.factory_id.data,
            department_id=department.id if department else None,
            date_of_joining=form.date_of_joining.data,
            email=(form.email.data or "").strip() or None,
            receive_alerts=receive_alerts,
            is_active=(form.is_active.data == "1"),
        )
        db.session.add(employee)
        db.session.commit()
        flash("Employee added successfully.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/form.html", form=form, title="Add Employee")


@employees_bp.route("/edit/<int:employee_id>", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def edit_employee(employee_id):
    employee = _get_employee_or_403(employee_id)
    form = EmployeeForm(obj=employee)
    _set_factory_choices(form)
    if request.method == "GET":
        form.is_active.data = "1" if employee.is_active else "0"
    if form.validate_on_submit():
        if not can_access_factory(form.factory_id.data):
            abort(403)
        dup = Employee.query.filter(
            Employee.factory_id == form.factory_id.data,
            Employee.employee_code == form.employee_code.data.strip(),
            Employee.id != employee.id,
        ).first()
        if dup:
            flash("Another employee with this ID already exists in this factory.", "danger")
            return render_template("employees/form.html", form=form, title="Edit Employee")
        department = _department_for_factory(form.department_id.data, form.factory_id.data)
        if department is False:
            flash("Choose a department from the list.", "danger")
            return render_template("employees/form.html", form=form, title="Edit Employee")
        receive_alerts = bool(form.receive_alerts.data) and valid_email(form.email.data)
        if form.receive_alerts.data and not receive_alerts:
            flash("Employee alerts stay off until a valid email is saved.", "warning")
        employee.employee_code = form.employee_code.data.strip()
        employee.name = form.name.data.strip()
        employee.designation = form.designation.data
        employee.section = form.section.data
        employee.factory_id = form.factory_id.data
        employee.department_id = department.id if department else None
        employee.date_of_joining = form.date_of_joining.data
        employee.email = (form.email.data or "").strip() or None
        employee.receive_alerts = receive_alerts
        employee.is_active = (form.is_active.data == "1")
        db.session.commit()
        flash("Employee updated successfully.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/form.html", form=form, title="Edit Employee")


@employees_bp.route("/<int:employee_id>")
@login_required
def view_employee(employee_id):
    employee = _get_employee_or_403(employee_id)
    licenses = [c for c in employee.compliances if c.is_active]
    checkups = [h for h in employee.health_checkups if h.is_active]
    return render_template(
        "employees/detail.html",
        employee=employee,
        licenses=sorted(licenses, key=lambda c: (c.expiry_date is None, c.expiry_date or date.max)),
        checkups=sorted(checkups, key=lambda h: (h.expiry_date is None, h.expiry_date or date.max)),
        can_manage=current_user.role.name in MANAGE_ROLES,
    )


# --- Employee licenses ------------------------------------------------------

@employees_bp.route("/<int:employee_id>/license/add", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def add_license(employee_id):
    employee = _get_employee_or_403(employee_id)
    form = EmployeeComplianceForm()
    eligible, ineligible, eligible_ids = _set_record_recipient_choices(form, employee.factory_id, employee.department_id)
    if form.validate_on_submit():
        no_expiry = bool(form.no_expiry.data)
        expiry_date = None if no_expiry else form.expiry_date.data
        if expiry_before_issue(form.issue_date.data, expiry_date):
            flash("Expiry date must be on or after the issue date.", "danger")
            return render_template("employees/license_form.html", form=form, employee=employee,
                                   title="Add License", suggestions=LICENSE_TYPE_SUGGESTIONS,
                                   eligible_recipients=eligible, ineligible_recipients=ineligible)
        record = EmployeeCompliance(
            employee_id=employee.id,
            compliance_type=form.compliance_type.data,
            registration_number=form.registration_number.data,
            category=form.category.data,
            issuing_authority=form.issuing_authority.data,
            issue_date=form.issue_date.data,
            expiry_date=expiry_date,
            no_expiry=no_expiry,
            status=compute_status(expiry_date, no_expiry),
            remarks=form.remarks.data,
            created_by=current_user.id,
        )
        db.session.add(record)
        db.session.flush()
        _apply_record_recipients(record, employee, form, eligible_ids)
        if form.file.data:
            try:
                _attach_license_file(record, employee, form.file.data)
            except ValueError as exc:
                flash(str(exc), "warning")
            except OSError:
                flash("License saved, but the file could not be stored.", "warning")
        db.session.commit()
        flash("License added.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/license_form.html", form=form, employee=employee,
                           title="Add License", suggestions=LICENSE_TYPE_SUGGESTIONS,
                           eligible_recipients=eligible, ineligible_recipients=ineligible)


@employees_bp.route("/license/<int:record_id>/edit", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def edit_license(record_id):
    record = EmployeeCompliance.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    form = EmployeeComplianceForm(obj=record)
    eligible, ineligible, eligible_ids = _set_record_recipient_choices(form, employee.factory_id, employee.department_id)
    if request.method == "GET":
        form.alert_recipients.data = [user.id for user in record.alert_recipients]
    if form.validate_on_submit():
        no_expiry = bool(form.no_expiry.data)
        expiry_date = None if no_expiry else form.expiry_date.data
        if expiry_before_issue(form.issue_date.data, expiry_date):
            flash("Expiry date must be on or after the issue date.", "danger")
            return render_template("employees/license_form.html", form=form, employee=employee,
                                   title="Edit License", suggestions=LICENSE_TYPE_SUGGESTIONS,
                                   eligible_recipients=eligible, ineligible_recipients=ineligible)
        record.compliance_type = form.compliance_type.data
        record.registration_number = form.registration_number.data
        record.category = form.category.data
        record.issuing_authority = form.issuing_authority.data
        record.issue_date = form.issue_date.data
        record.expiry_date = expiry_date
        record.no_expiry = no_expiry
        record.status = compute_status(expiry_date, no_expiry)
        record.remarks = form.remarks.data
        _apply_record_recipients(record, employee, form, eligible_ids)
        if form.file.data:
            try:
                _attach_license_file(record, employee, form.file.data)
            except ValueError as exc:
                flash(str(exc), "warning")
            except OSError:
                flash("Saved, but the file could not be stored.", "warning")
        db.session.commit()
        flash("License updated.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/license_form.html", form=form, employee=employee,
                           title="Edit License", suggestions=LICENSE_TYPE_SUGGESTIONS,
                           eligible_recipients=eligible, ineligible_recipients=ineligible)


@employees_bp.route("/license/<int:record_id>/archive", methods=["POST"])
@login_required
@roles_required(*DELETE_ROLES)
def archive_license(record_id):
    record = EmployeeCompliance.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    record.is_active = False
    db.session.commit()
    flash("License archived.", "success")
    return redirect(url_for("employees.view_employee", employee_id=employee.id))


@employees_bp.route("/license/<int:record_id>/download")
@login_required
def download_license(record_id):
    record = EmployeeCompliance.query.get_or_404(record_id)
    _get_employee_or_403(record.employee_id)
    if not record.file_path:
        abort(404)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], record.file_path, as_attachment=True)


@employees_bp.route("/license/<int:record_id>/file")
@login_required
def view_license_file(record_id):
    record = EmployeeCompliance.query.get_or_404(record_id)
    _get_employee_or_403(record.employee_id)
    if not record.file_path:
        abort(404)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], record.file_path, as_attachment=False)


@employees_bp.route("/license/version/<int:version_id>/<string:mode>")
@login_required
def license_version_file(version_id, mode):
    version = EmployeeComplianceVersion.query.get_or_404(version_id)
    record = EmployeeCompliance.query.get_or_404(version.compliance_id)
    _get_employee_or_403(record.employee_id)
    if mode not in ("view", "download"):
        abort(404)
    return send_from_directory(
        current_app.config["UPLOAD_FOLDER"], version.file_path, as_attachment=(mode == "download"),
    )


# --- Health checkups --------------------------------------------------------

@employees_bp.route("/<int:employee_id>/checkup/add", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def add_checkup(employee_id):
    employee = _get_employee_or_403(employee_id)
    form = HealthCheckupForm()
    eligible, ineligible, eligible_ids = _set_record_recipient_choices(form, employee.factory_id, employee.department_id)
    if form.validate_on_submit():
        no_expiry = bool(form.no_expiry.data)
        expiry_date = None if no_expiry else form.expiry_date.data
        if expiry_before_issue(form.report_date.data, expiry_date):
            flash("Expiry date must be on or after the report date.", "danger")
            return render_template("employees/checkup_form.html", form=form, employee=employee, title="Add Health Checkup",
                                   eligible_recipients=eligible, ineligible_recipients=ineligible)
        record = HealthCheckup(
            employee_id=employee.id,
            exam_type=form.exam_type.data,
            report_date=form.report_date.data,
            expiry_date=expiry_date,
            no_expiry=no_expiry,
            status=compute_status(expiry_date, no_expiry),
            pre_yn=form.pre_yn.data or None,
            test_center=form.test_center.data,
            remarks=form.remarks.data,
            created_by=current_user.id,
        )
        db.session.add(record)
        db.session.flush()
        _apply_record_recipients(record, employee, form, eligible_ids)
        if form.file.data:
            try:
                _attach_checkup_file(record, employee, form.file.data)
            except ValueError as exc:
                flash(str(exc), "warning")
            except OSError:
                flash("Checkup saved, but the file could not be stored.", "warning")
        db.session.commit()
        flash("Health checkup added.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/checkup_form.html", form=form, employee=employee, title="Add Health Checkup",
                           eligible_recipients=eligible, ineligible_recipients=ineligible)


@employees_bp.route("/checkup/<int:record_id>/edit", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def edit_checkup(record_id):
    record = HealthCheckup.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    form = HealthCheckupForm(obj=record)
    eligible, ineligible, eligible_ids = _set_record_recipient_choices(form, employee.factory_id, employee.department_id)
    if request.method == "GET":
        form.alert_recipients.data = [user.id for user in record.alert_recipients]
    if form.validate_on_submit():
        no_expiry = bool(form.no_expiry.data)
        expiry_date = None if no_expiry else form.expiry_date.data
        if expiry_before_issue(form.report_date.data, expiry_date):
            flash("Expiry date must be on or after the report date.", "danger")
            return render_template("employees/checkup_form.html", form=form, employee=employee, title="Edit Health Checkup",
                                   eligible_recipients=eligible, ineligible_recipients=ineligible)
        record.exam_type = form.exam_type.data
        record.report_date = form.report_date.data
        record.expiry_date = expiry_date
        record.no_expiry = no_expiry
        record.status = compute_status(expiry_date, no_expiry)
        record.pre_yn = form.pre_yn.data or None
        record.test_center = form.test_center.data
        record.remarks = form.remarks.data
        _apply_record_recipients(record, employee, form, eligible_ids)
        if form.file.data:
            try:
                _attach_checkup_file(record, employee, form.file.data)
            except ValueError as exc:
                flash(str(exc), "warning")
            except OSError:
                flash("Saved, but the file could not be stored.", "warning")
        db.session.commit()
        flash("Health checkup updated.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/checkup_form.html", form=form, employee=employee, title="Edit Health Checkup",
                           eligible_recipients=eligible, ineligible_recipients=ineligible)


@employees_bp.route("/checkup/<int:record_id>/archive", methods=["POST"])
@login_required
@roles_required(*DELETE_ROLES)
def archive_checkup(record_id):
    record = HealthCheckup.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    record.is_active = False
    db.session.commit()
    flash("Health checkup archived.", "success")
    return redirect(url_for("employees.view_employee", employee_id=employee.id))


@employees_bp.route("/checkup/<int:record_id>/download")
@login_required
def download_checkup(record_id):
    record = HealthCheckup.query.get_or_404(record_id)
    _get_employee_or_403(record.employee_id)
    if not record.file_path:
        abort(404)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], record.file_path, as_attachment=True)


@employees_bp.route("/checkup/<int:record_id>/file")
@login_required
def view_checkup_file(record_id):
    record = HealthCheckup.query.get_or_404(record_id)
    _get_employee_or_403(record.employee_id)
    if not record.file_path:
        abort(404)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], record.file_path, as_attachment=False)


@employees_bp.route("/checkup/version/<int:version_id>/<string:mode>")
@login_required
def checkup_version_file(version_id, mode):
    version = HealthCheckupVersion.query.get_or_404(version_id)
    record = HealthCheckup.query.get_or_404(version.checkup_id)
    _get_employee_or_403(record.employee_id)
    if mode not in ("view", "download"):
        abort(404)
    return send_from_directory(
        current_app.config["UPLOAD_FOLDER"], version.file_path, as_attachment=(mode == "download"),
    )


# --- Cross-employee compliance / checkup lists (search + filter) ------------

@employees_bp.route("/licenses")
@login_required
def list_licenses():
    query = EmployeeCompliance.query.join(Employee).options(joinedload(EmployeeCompliance.employee)) \
        .filter(EmployeeCompliance.is_active.is_(True))
    if not current_user.can_access_all_factories():
        query = query.filter(Employee.factory_id == current_user.factory_id) if current_user.factory_id else query.filter(db.false())

    compliance_type = request.args.get("compliance_type", "").strip()
    status = request.args.get("status", "")
    search = request.args.get("search", "").strip()
    if compliance_type:
        query = query.filter(EmployeeCompliance.compliance_type.ilike(f"%{compliance_type}%"))
    if status:
        query = query.filter(EmployeeCompliance.status == status)
    if search:
        from app.ocr_service import normalize_bangla_digits
        like = f"%{search}%"
        norm = f"%{normalize_bangla_digits(search)}%"
        ocr_ids = db.session.query(OcrResult.employee_compliance_id).filter(
            OcrResult.employee_compliance_id.isnot(None),
            db.or_(OcrResult.extracted_text.ilike(like), OcrResult.search_text.ilike(norm)),
        )
        query = query.filter(db.or_(
            EmployeeCompliance.registration_number.ilike(like),
            Employee.name.ilike(like),
            EmployeeCompliance.id.in_(ocr_ids),
        ))

    page = request.args.get("page", 1, type=int)
    pagination = query.order_by(EmployeeCompliance.expiry_date.is_(None), EmployeeCompliance.expiry_date) \
        .paginate(page=page, per_page=15, error_out=False)
    return render_template("employees/licenses.html", records=pagination.items, pagination=pagination,
                           selected_type=compliance_type, selected_status=status, search=search,
                           all_statuses=ALL_STATUSES)


@employees_bp.route("/checkups")
@login_required
def list_checkups():
    query = HealthCheckup.query.join(Employee).options(joinedload(HealthCheckup.employee)) \
        .filter(HealthCheckup.is_active.is_(True))
    if not current_user.can_access_all_factories():
        query = query.filter(Employee.factory_id == current_user.factory_id) if current_user.factory_id else query.filter(db.false())

    exam_type = request.args.get("exam_type", "").strip()
    status = request.args.get("status", "")
    search = request.args.get("search", "").strip()
    if exam_type:
        query = query.filter(HealthCheckup.exam_type == exam_type)
    if status:
        query = query.filter(HealthCheckup.status == status)
    if search:
        from app.ocr_service import normalize_bangla_digits
        like = f"%{search}%"
        norm = f"%{normalize_bangla_digits(search)}%"
        ocr_ids = db.session.query(OcrResult.health_checkup_id).filter(
            OcrResult.health_checkup_id.isnot(None),
            db.or_(OcrResult.extracted_text.ilike(like), OcrResult.search_text.ilike(norm)),
        )
        query = query.filter(db.or_(
            HealthCheckup.exam_type.ilike(like),
            Employee.name.ilike(like),
            HealthCheckup.id.in_(ocr_ids),
        ))

    page = request.args.get("page", 1, type=int)
    pagination = query.order_by(HealthCheckup.expiry_date.is_(None), HealthCheckup.expiry_date) \
        .paginate(page=page, per_page=15, error_out=False)
    return render_template("employees/checkups.html", records=pagination.items, pagination=pagination,
                           selected_exam=exam_type, selected_status=status, search=search,
                           all_statuses=ALL_STATUSES,
                           exam_types=["CBC", "Hearing / Widal", "Chest X-Ray P/A", "Urine Test"])


# --- Excel import (Super Admin only) ----------------------------------------

def _imports_dir():
    path = os.path.join(current_app.config["UPLOAD_FOLDER"], "_imports")
    os.makedirs(path, exist_ok=True)
    return path


@employees_bp.route("/import", methods=["GET", "POST"])
@login_required
@roles_required("Super Admin")
def import_data():
    from app.import_utils import IMPORT_TYPES, parse_workbook, build_factory_lookup

    if request.method == "POST":
        import_type = request.form.get("import_type", "")
        file = request.files.get("file")
        if import_type not in IMPORT_TYPES:
            flash("Please choose a valid import type.", "danger")
            return redirect(url_for("employees.import_data"))
        if not file or not file.filename:
            flash("Please choose an Excel (.xlsx) file.", "danger")
            return redirect(url_for("employees.import_data"))
        if not file.filename.lower().endswith(".xlsx"):
            flash("Only .xlsx files are supported for import.", "danger")
            return redirect(url_for("employees.import_data"))

        token = secrets.token_hex(16)
        temp_path = os.path.join(_imports_dir(), f"{token}.xlsx")
        file.save(temp_path)

        try:
            from openpyxl import load_workbook
            from app.import_utils import annotate_existing, parse_full_workbook, preview_counts
            wb = load_workbook(temp_path, read_only=True, data_only=True)
            if import_type == "workbook":
                parsed, errors, warnings = parse_full_workbook(wb, build_factory_lookup())
            else:
                parsed, errors, warnings = parse_workbook(wb.active, import_type, build_factory_lookup())
            wb.close()
            importable, duplicate_warnings = annotate_existing(parsed)
            warnings = list(warnings) + duplicate_warnings
            counts = preview_counts(importable, errors, warnings)
        except Exception:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            log_event("XLSX_IMPORT_FAILED", f"Excel preview ({import_type}) failed.", level="ERROR")
            flash("Could not read the Excel file. Use a valid .xlsx workbook.", "danger")
            return redirect(url_for("employees.import_data"))

        return render_template("employees/import_preview.html",
                               import_type=import_type, import_type_label=IMPORT_TYPES[import_type],
                               parsed=importable, errors=errors, warnings=warnings, token=token,
                               counts=counts, preview_rows=importable[:50])

    return render_template("employees/import.html", import_types=IMPORT_TYPES)


@employees_bp.route("/import/confirm", methods=["POST"])
@login_required
@roles_required("Super Admin")
def import_confirm():
    from app.import_utils import (
        IMPORT_TYPES, annotate_existing, apply_import, build_factory_lookup,
        parse_full_workbook, parse_workbook, preview_counts,
    )

    import_type = request.form.get("import_type", "")
    token = request.form.get("token", "")
    if import_type not in IMPORT_TYPES or not token or not token.isalnum():
        flash("Invalid import request.", "danger")
        return redirect(url_for("employees.import_data"))

    temp_path = os.path.join(_imports_dir(), f"{token}.xlsx")
    if not os.path.exists(temp_path):
        flash("The uploaded file expired. Please upload again.", "warning")
        return redirect(url_for("employees.import_data"))

    log_event("XLSX_IMPORT_STARTED", f"Excel import started ({IMPORT_TYPES.get(import_type, import_type)}).")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(temp_path, read_only=True, data_only=True)
        if import_type == "workbook":
            parsed, errors, warnings = parse_full_workbook(wb, build_factory_lookup())
        else:
            parsed, errors, warnings = parse_workbook(wb.active, import_type, build_factory_lookup())
        wb.close()
        importable, duplicate_warnings = annotate_existing(parsed)
        warnings = list(warnings) + duplicate_warnings
        counts = preview_counts(importable, errors, warnings)
        summary = apply_import(importable)
    except Exception:
        db.session.rollback()
        log_event("XLSX_IMPORT_FAILED", f"Excel import ({import_type}) failed.", level="ERROR")
        flash("Import failed. The file was not applied. Check the file and try again.", "danger")
        return redirect(url_for("employees.import_data"))
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

    imported = (
        summary["employees_created"] + summary["licenses_created"]
        + summary["checkups_created"] + summary["recipients_linked"]
    )
    log_event(
        "XLSX_IMPORT_COMPLETED",
        f"Excel import ({IMPORT_TYPES[import_type]}): total {counts['total']}, "
        f"valid {counts['valid']}, warning {counts['warning']}, failed {counts['failed']}, "
        f"imported {imported}.",
    )
    flash(
        f"Import summary — Total rows: {counts['total']}. Valid: {counts['valid']}. "
        f"Warning: {counts['warning']}. Failed: {counts['failed']}. Imported: {imported}.",
        "success" if counts["failed"] == 0 else "warning",
    )
    return redirect(url_for("employees.list_employees"))


def _safe_employee_file(relative_path):
    root = os.path.realpath(current_app.config["UPLOAD_FOLDER"])
    full = os.path.realpath(os.path.join(root, (relative_path or "").replace("\\", "/")))
    if os.path.commonpath([root, full]) != root or not os.path.isfile(full):
        abort(404)
    return full


def _store_record_ocr(record, kind):
    from datetime import datetime
    from app.ocr_service import extract_document, OcrUnavailable, OcrError, suggestions_json
    if not record.file_path:
        raise OcrError("Upload a PDF, JPG, or PNG before running OCR.")
    ext = record.file_path.rsplit(".", 1)[-1].lower()
    if ext not in {"pdf", "jpg", "jpeg", "png"}:
        raise OcrError("OCR supports PDF, JPG, and PNG files only.")
    label = "License" if kind == "license" else "Checkup"
    log_event("OCR_STARTED", f"{label} ID: {record.id}.")
    try:
        result = extract_document(_safe_employee_file(record.file_path), current_app.config)
    except OcrUnavailable:
        log_event("OCR_FAILED", f"{label} ID: {record.id}. Engine unavailable.", level="ERROR")
        raise
    except OcrError as exc:
        row = record.ocr_result or OcrResult(**{("employee_compliance_id" if kind == "license" else "health_checkup_id"): record.id})
        row.ocr_status = "FAILED"
        row.error_message = str(exc)[:400]
        row.processed_at = datetime.utcnow()
        db.session.add(row)
        db.session.commit()
        log_event("OCR_FAILED", f"{label} ID: {record.id}. Failed.", level="ERROR")
        raise
    key = "employee_compliance_id" if kind == "license" else "health_checkup_id"
    row = record.ocr_result or OcrResult(**{key: record.id})
    row.extracted_text = result["extracted_text"]
    row.search_text = result["search_text"]
    row.language = result["language"]
    row.confidence = result["confidence"]
    row.suggested_fields = suggestions_json(result["suggested_fields"])
    row.ocr_status = "COMPLETED"
    row.error_message = None
    row.confirmed_at = None
    row.processed_at = result["processed_at"]
    row.processing_time = result["processing_time"]
    db.session.add(row)
    db.session.commit()
    log_event("OCR_COMPLETED", f"{label} ID: {record.id}. Language: {result['language']}.")
    return row


def _confirm_record_ocr(record, kind):
    from datetime import datetime
    from app.ocr_service import normalize_bangla_digits, parse_confirmed_date, load_suggestions
    row = record.ocr_result
    if not row or row.ocr_status != "COMPLETED":
        flash("There is no completed OCR result to confirm.", "warning")
        return
    edited = request.form.get("extracted_text", "")
    row.extracted_text = edited
    row.search_text = normalize_bangla_digits(edited)
    skipped = []
    if kind == "license" and request.form.get("apply_number"):
        number = (request.form.get("registration_number") or "").strip()
        if number:
            record.registration_number = number
    if kind == "checkup" and request.form.get("apply_center"):
        center = (request.form.get("test_center") or "").strip()
        if center:
            record.test_center = center
    for field in ("issue_date", "expiry_date", "report_date"):
        if not request.form.get(f"apply_{field}"):
            continue
        if not hasattr(record, field):
            continue
        parsed = parse_confirmed_date(request.form.get(field) or "")
        if parsed is None:
            skipped.append(field.replace("_", " "))
            continue
        setattr(record, field, parsed)
        if field == "expiry_date":
            record.no_expiry = False
            record.status = compute_status(parsed, False)
    row.confirmed_at = datetime.utcnow()
    db.session.commit()
    flash("Confirmed OCR text saved.", "success")
    if skipped:
        flash("These dates were not applied because the format is ambiguous: " + ", ".join(skipped) + ".", "warning")
    return load_suggestions(row.suggested_fields)


@employees_bp.route("/license/<int:record_id>/ocr", methods=["POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def run_license_ocr(record_id):
    from app.ocr_service import OcrUnavailable, OcrError
    record = EmployeeCompliance.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    try:
        _store_record_ocr(record, "license")
    except OcrUnavailable as exc:
        flash(str(exc), "warning")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    except OcrError as exc:
        flash(str(exc), "danger")
    return redirect(url_for("employees.view_license_ocr", record_id=record.id))


@employees_bp.route("/license/<int:record_id>/ocr/view", methods=["GET", "POST"])
@login_required
def view_license_ocr(record_id):
    from app.ocr_service import load_suggestions
    record = EmployeeCompliance.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    can_manage = current_user.role.name in MANAGE_ROLES
    if request.method == "POST":
        if not can_manage:
            abort(403)
        _confirm_record_ocr(record, "license")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/ocr.html", employee=employee, record=record, kind="license",
                           ocr=record.ocr_result,
                           suggestions=load_suggestions(record.ocr_result.suggested_fields) if record.ocr_result else {},
                           can_manage=can_manage)


@employees_bp.route("/checkup/<int:record_id>/ocr", methods=["POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def run_checkup_ocr(record_id):
    from app.ocr_service import OcrUnavailable, OcrError
    record = HealthCheckup.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    try:
        _store_record_ocr(record, "checkup")
    except OcrUnavailable as exc:
        flash(str(exc), "warning")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    except OcrError as exc:
        flash(str(exc), "danger")
    return redirect(url_for("employees.view_checkup_ocr", record_id=record.id))


@employees_bp.route("/checkup/<int:record_id>/ocr/view", methods=["GET", "POST"])
@login_required
def view_checkup_ocr(record_id):
    from app.ocr_service import load_suggestions
    record = HealthCheckup.query.get_or_404(record_id)
    employee = _get_employee_or_403(record.employee_id)
    can_manage = current_user.role.name in MANAGE_ROLES
    if request.method == "POST":
        if not can_manage:
            abort(403)
        _confirm_record_ocr(record, "checkup")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template("employees/ocr.html", employee=employee, record=record, kind="checkup",
                           ocr=record.ocr_result,
                           suggestions=load_suggestions(record.ocr_result.suggested_fields) if record.ocr_result else {},
                           can_manage=can_manage)


# --- Pre-save OCR extract preview ------------------------------------------
#
# The flows below mirror documents.extract_upload / extract_preview but target
# employee licenses and health checkups: upload a PDF/JPG/PNG, run OCR into a
# staging folder, redirect to a confirmation form that is pre-filled from the
# suggested fields. On confirm, the record, its first file version, and a
# COMPLETED OcrResult are all written in one go. OCR is a suggestion engine;
# the user still owns the final field values.

_OCR_EXTS = {"pdf", "jpg", "jpeg", "png"}


def _employee_staging_dir():
    path = os.path.join(current_app.instance_path, "ocr_staging")
    os.makedirs(path, exist_ok=True)
    return path


def _load_employee_staging(token, expected_kind):
    import json
    if not token or not token.isalnum():
        abort(404)
    folder = _employee_staging_dir()
    meta_path = os.path.join(folder, f"{token}.json")
    if not os.path.isfile(meta_path):
        abort(404)
    with open(meta_path, encoding="utf-8") as handle:
        meta = json.load(handle)
    if meta.get("kind") != expected_kind:
        abort(404)
    if meta.get("user_id") != current_user.id and not current_user.can_access_all_factories():
        abort(403)
    return meta, meta_path


def _stage_employee_upload(employee, kind, redirect_url):
    """Shared implementation of the extract-upload POST for license + checkup."""
    import json
    from app.ocr_service import extract_document, OcrUnavailable, OcrError
    upload = request.files.get("file")
    if not upload or not upload.filename:
        flash("Choose a PDF, JPG, or PNG file to extract.", "warning")
        return redirect(redirect_url)
    original = upload.filename
    ext = original.rsplit(".", 1)[-1].lower() if "." in original else ""
    if ext not in _OCR_EXTS:
        flash("OCR supports PDF, JPG, and PNG files only.", "warning")
        return redirect(redirect_url)
    token = secrets.token_hex(16)
    folder = _employee_staging_dir()
    stored = os.path.join(folder, f"{token}.{ext}")
    upload.save(stored)
    log_event("OCR_STARTED", f"{kind.capitalize()} upload extraction. Employee ID: {employee.id}.")
    try:
        result = extract_document(stored, current_app.config)
    except (OcrUnavailable, OcrError) as exc:
        if os.path.exists(stored):
            os.remove(stored)
        log_event("OCR_FAILED", f"{kind.capitalize()} upload extraction failed.", level="ERROR")
        flash(str(exc), "warning" if isinstance(exc, OcrUnavailable) else "danger")
        return redirect(redirect_url)
    meta = {
        "kind": kind,
        "employee_id": employee.id,
        "user_id": current_user.id,
        "original_name": original,
        "result": {
            "extracted_text": result["extracted_text"],
            "language": result["language"],
            "confidence": result["confidence"],
            "suggested_fields": result["suggested_fields"],
            "processing_time": result["processing_time"],
        },
    }
    with open(os.path.join(folder, f"{token}.json"), "w", encoding="utf-8") as handle:
        json.dump(meta, handle, ensure_ascii=False)
    log_event("OCR_COMPLETED",
              f"{kind.capitalize()} upload extraction. Employee ID: {employee.id}. "
              f"Language: {result['language']}.")
    endpoint = ("employees.extract_license_preview" if kind == "license"
                else "employees.extract_checkup_preview")
    return redirect(url_for(endpoint, employee_id=employee.id, token=token))


def _staged_file_for(token):
    """Return the staged upload path for a token, or None if it has been removed."""
    folder = _employee_staging_dir()
    for name in os.listdir(folder):
        if name.startswith(token + ".") and not name.endswith(".json"):
            return os.path.join(folder, name)
    return None


def _cancel_staging(token):
    folder = _employee_staging_dir()
    for name in list(os.listdir(folder)):
        if name == f"{token}.json" or (name.startswith(token + ".") and not name.endswith(".json")):
            try:
                os.remove(os.path.join(folder, name))
            except OSError:
                pass


def _fill_license_form(form, suggestions):
    from datetime import date as _date
    suggestions = suggestions or {}
    compliance_type = (suggestions.get("doc_type")
                       or suggestions.get("doc_name")
                       or "").strip()
    if compliance_type:
        canonical = next((s for s in LICENSE_TYPE_SUGGESTIONS
                          if s.lower() == compliance_type.lower()), compliance_type)
        form.compliance_type.data = canonical[:100]
    reg = (suggestions.get("registration_number")
           or suggestions.get("license_no")
           or suggestions.get("cert_no")
           or "")
    form.registration_number.data = (reg or "")[:100]
    form.issuing_authority.data = (suggestions.get("issuing_authority") or "")[:200]
    for field_name in ("issue_date", "expiry_date"):
        hint = suggestions.get(field_name) or {}
        if hint.get("iso") and not hint.get("ambiguous"):
            getattr(form, field_name).data = _date.fromisoformat(hint["iso"])


_CHECKUP_TYPE_CHOICES = ("CBC", "Hearing / Widal", "Chest X-Ray P/A", "Urine Test")


def _match_checkup_type(value):
    needle = (value or "").strip().lower()
    if not needle:
        return None
    for choice in _CHECKUP_TYPE_CHOICES:
        if choice.lower() == needle or choice.lower() in needle or needle in choice.lower():
            return choice
    return None


def _fill_checkup_form(form, suggestions):
    from datetime import date as _date
    suggestions = suggestions or {}
    exam = _match_checkup_type(suggestions.get("doc_type")) \
        or _match_checkup_type(suggestions.get("doc_name"))
    if exam:
        form.exam_type.data = exam
    form.test_center.data = (suggestions.get("test_center") or "")[:200]
    report_hint = suggestions.get("report_date") or suggestions.get("issue_date") or {}
    if report_hint.get("iso") and not report_hint.get("ambiguous"):
        form.report_date.data = _date.fromisoformat(report_hint["iso"])
    expiry_hint = suggestions.get("expiry_date") or {}
    if expiry_hint.get("iso") and not expiry_hint.get("ambiguous"):
        form.expiry_date.data = _date.fromisoformat(expiry_hint["iso"])


def _persist_extracted_record(employee, record, kind, meta, suggestions):
    """Attach the staged upload to the saved record and write a COMPLETED OcrResult."""
    from datetime import datetime
    from werkzeug.datastructures import FileStorage
    from app.ocr_service import normalize_bangla_digits, suggestions_json
    token = meta["token"]
    staged = _staged_file_for(token)
    if staged and os.path.exists(staged):
        original_name = meta.get("original_name") or os.path.basename(staged)
        with open(staged, "rb") as handle:
            storage = FileStorage(handle, filename=original_name)
            if kind == "license":
                _attach_license_file(record, employee, storage)
            else:
                _attach_checkup_file(record, employee, storage)
    db.session.flush()
    text = request.form.get("extracted_text", meta["result"].get("extracted_text") or "")
    key = "employee_compliance_id" if kind == "license" else "health_checkup_id"
    db.session.add(OcrResult(**{
        key: record.id,
        "extracted_text": text,
        "search_text": normalize_bangla_digits(text),
        "language": meta["result"].get("language"),
        "confidence": meta["result"].get("confidence"),
        "suggested_fields": suggestions_json(suggestions),
        "ocr_status": "COMPLETED",
        "confirmed_at": datetime.utcnow(),
        "processed_at": datetime.utcnow(),
        "processing_time": meta["result"].get("processing_time"),
    }))


@employees_bp.route("/<int:employee_id>/license/extract", methods=["POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def extract_license_upload(employee_id):
    employee = _get_employee_or_403(employee_id)
    return _stage_employee_upload(
        employee, "license",
        url_for("employees.add_license", employee_id=employee.id),
    )


@employees_bp.route("/<int:employee_id>/license/extract/<token>", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def extract_license_preview(employee_id, token):
    employee = _get_employee_or_403(employee_id)
    meta, meta_path = _load_employee_staging(token, "license")
    if meta.get("employee_id") != employee.id:
        abort(404)
    suggestions = meta["result"].get("suggested_fields") or {}
    form = EmployeeComplianceForm()
    eligible, ineligible, eligible_ids = _set_record_recipient_choices(
        form, employee.factory_id, employee.department_id,
    )
    if request.method == "GET":
        _fill_license_form(form, suggestions)
    if request.method == "POST" and request.form.get("action") == "cancel":
        _cancel_staging(token)
        flash("Extraction cancelled. Nothing was saved.", "info")
        return redirect(url_for("employees.add_license", employee_id=employee.id))
    if request.method == "POST":
        form.file.data = None  # file comes from staging, not re-upload
    if form.validate_on_submit():
        no_expiry = bool(form.no_expiry.data)
        expiry_date = None if no_expiry else form.expiry_date.data
        if expiry_before_issue(form.issue_date.data, expiry_date):
            flash("Expiry date must be on or after the issue date.", "danger")
            return render_template(
                "employees/extract_license_preview.html",
                form=form, employee=employee, token=token,
                result=meta["result"], suggestions=suggestions,
                type_suggestions=LICENSE_TYPE_SUGGESTIONS,
                eligible_recipients=eligible, ineligible_recipients=ineligible,
            )
        record = EmployeeCompliance(
            employee_id=employee.id,
            compliance_type=form.compliance_type.data,
            registration_number=form.registration_number.data,
            category=form.category.data,
            issuing_authority=form.issuing_authority.data,
            issue_date=form.issue_date.data,
            expiry_date=expiry_date,
            no_expiry=no_expiry,
            status=compute_status(expiry_date, no_expiry),
            remarks=form.remarks.data,
            created_by=current_user.id,
        )
        db.session.add(record)
        db.session.flush()
        _apply_record_recipients(record, employee, form, eligible_ids)
        meta["token"] = token
        _persist_extracted_record(employee, record, "license", meta, suggestions)
        db.session.commit()
        _cancel_staging(token)
        if os.path.exists(meta_path):
            try:
                os.remove(meta_path)
            except OSError:
                pass
        flash("License saved from the confirmed extraction.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template(
        "employees/extract_license_preview.html",
        form=form, employee=employee, token=token,
        result=meta["result"], suggestions=suggestions,
        type_suggestions=LICENSE_TYPE_SUGGESTIONS,
        eligible_recipients=eligible, ineligible_recipients=ineligible,
    )


@employees_bp.route("/<int:employee_id>/checkup/extract", methods=["POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def extract_checkup_upload(employee_id):
    employee = _get_employee_or_403(employee_id)
    return _stage_employee_upload(
        employee, "checkup",
        url_for("employees.add_checkup", employee_id=employee.id),
    )


@employees_bp.route("/<int:employee_id>/checkup/extract/<token>", methods=["GET", "POST"])
@login_required
@roles_required(*MANAGE_ROLES)
def extract_checkup_preview(employee_id, token):
    employee = _get_employee_or_403(employee_id)
    meta, meta_path = _load_employee_staging(token, "checkup")
    if meta.get("employee_id") != employee.id:
        abort(404)
    suggestions = meta["result"].get("suggested_fields") or {}
    form = HealthCheckupForm()
    eligible, ineligible, eligible_ids = _set_record_recipient_choices(
        form, employee.factory_id, employee.department_id,
    )
    if request.method == "GET":
        _fill_checkup_form(form, suggestions)
    if request.method == "POST" and request.form.get("action") == "cancel":
        _cancel_staging(token)
        flash("Extraction cancelled. Nothing was saved.", "info")
        return redirect(url_for("employees.add_checkup", employee_id=employee.id))
    if request.method == "POST":
        form.file.data = None
    if form.validate_on_submit():
        no_expiry = bool(form.no_expiry.data)
        expiry_date = None if no_expiry else form.expiry_date.data
        if expiry_before_issue(form.report_date.data, expiry_date):
            flash("Expiry date must be on or after the report date.", "danger")
            return render_template(
                "employees/extract_checkup_preview.html",
                form=form, employee=employee, token=token,
                result=meta["result"], suggestions=suggestions,
                eligible_recipients=eligible, ineligible_recipients=ineligible,
            )
        record = HealthCheckup(
            employee_id=employee.id,
            exam_type=form.exam_type.data,
            report_date=form.report_date.data,
            expiry_date=expiry_date,
            no_expiry=no_expiry,
            status=compute_status(expiry_date, no_expiry),
            pre_yn=form.pre_yn.data or None,
            test_center=form.test_center.data,
            remarks=form.remarks.data,
            created_by=current_user.id,
        )
        db.session.add(record)
        db.session.flush()
        _apply_record_recipients(record, employee, form, eligible_ids)
        meta["token"] = token
        _persist_extracted_record(employee, record, "checkup", meta, suggestions)
        db.session.commit()
        _cancel_staging(token)
        if os.path.exists(meta_path):
            try:
                os.remove(meta_path)
            except OSError:
                pass
        flash("Health checkup saved from the confirmed extraction.", "success")
        return redirect(url_for("employees.view_employee", employee_id=employee.id))
    return render_template(
        "employees/extract_checkup_preview.html",
        form=form, employee=employee, token=token,
        result=meta["result"], suggestions=suggestions,
        eligible_recipients=eligible, ineligible_recipients=ineligible,
    )
